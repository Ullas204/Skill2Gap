"""Synthetic resume generator.

Builds realistic resume content from the data pool and renders it to
PDF (reportlab), DOCX (python-docx) or plain text so that generated
resumes round-trip through the existing ResumeParser.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field

from app.services.demo import data_pool


@dataclass
class ResumeEducation:
    degree: str
    institution: str
    start: str
    end: str
    gpa: str | None = None


@dataclass
class ResumeExperience:
    title: str
    company: str
    start: str
    end: str
    bullets: list[str] = field(default_factory=list)


@dataclass
class ResumeProject:
    name: str
    technologies: list[str]
    bullets: list[str] = field(default_factory=list)


@dataclass
class ResumeContent:
    full_name: str
    title: str
    email: str
    phone: str
    location: str
    summary: str
    skills: list[str]
    education: list[ResumeEducation] = field(default_factory=list)
    experiences: list[ResumeExperience] = field(default_factory=list)
    projects: list[ResumeProject] = field(default_factory=list)
    certifications: list[str] = field(default_factory=list)
    languages: list[dict] = field(default_factory=list)

    def to_text(self) -> str:
        lines: list[str] = []
        lines.append(self.full_name)
        lines.append(self.title)
        lines.append(f"{self.email} | {self.phone} | {self.location}")
        lines.append("")

        lines.append("PROFESSIONAL SUMMARY")
        lines.append(self.summary)
        lines.append("")

        lines.append("SKILLS")
        lines.append(", ".join(self.skills))
        lines.append("")

        if self.experiences:
            lines.append("PROFESSIONAL EXPERIENCE")
            for exp in self.experiences:
                lines.append(f"{exp.title} at {exp.company}")
                lines.append(f"{exp.start} - {exp.end}")
                for bullet in exp.bullets:
                    lines.append(f"- {bullet}")
                lines.append("")

        if self.education:
            lines.append("EDUCATION")
            for edu in self.education:
                gpa = f" | GPA: {edu.gpa}" if edu.gpa else ""
                lines.append(f"{edu.degree} - {edu.institution}")
                lines.append(f"{edu.start} - {edu.end}{gpa}")
            lines.append("")

        if self.projects:
            lines.append("PROJECTS")
            for proj in self.projects:
                lines.append(proj.name)
                lines.append(f"Technologies: {', '.join(proj.technologies)}")
                for bullet in proj.bullets:
                    lines.append(f"- {bullet}")
                lines.append("")

        if self.certifications:
            lines.append("CERTIFICATIONS")
            for cert in self.certifications:
                lines.append(cert)
            lines.append("")

        if self.languages:
            lines.append("LANGUAGES")
            for lang in self.languages:
                lines.append(f"{lang['name']} - {lang['proficiency']}")
            lines.append("")

        return "\n".join(lines).strip()

    def to_docx_bytes(self) -> bytes:
        from docx import Document

        doc = Document()
        doc.add_heading(self.full_name, level=0)
        doc.add_paragraph(self.title)
        doc.add_paragraph(f"{self.email} | {self.phone} | {self.location}")
        doc.add_paragraph("")

        doc.add_heading("Professional Summary", level=1)
        doc.add_paragraph(self.summary)

        doc.add_heading("Skills", level=1)
        doc.add_paragraph(", ".join(self.skills))

        if self.experiences:
            doc.add_heading("Professional Experience", level=1)
            for exp in self.experiences:
                doc.add_paragraph(f"{exp.title} at {exp.company}  |  {exp.start} - {exp.end}")
                for bullet in exp.bullets:
                    doc.add_paragraph(f"- {bullet}", style="List Bullet")

        if self.education:
            doc.add_heading("Education", level=1)
            for edu in self.education:
                gpa = f" | GPA: {edu.gpa}" if edu.gpa else ""
                doc.add_paragraph(f"{edu.degree} - {edu.institution}  |  {edu.start} - {edu.end}{gpa}")

        if self.projects:
            doc.add_heading("Projects", level=1)
            for proj in self.projects:
                doc.add_paragraph(proj.name)
                doc.add_paragraph(f"Technologies: {', '.join(proj.technologies)}")
                for bullet in proj.bullets:
                    doc.add_paragraph(f"- {bullet}", style="List Bullet")

        if self.certifications:
            doc.add_heading("Certifications", level=1)
            for cert in self.certifications:
                doc.add_paragraph(cert)

        if self.languages:
            doc.add_heading("Languages", level=1)
            for lang in self.languages:
                doc.add_paragraph(f"{lang['name']} - {lang['proficiency']}")

        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()

    def to_pdf_bytes(self) -> bytes:
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import inch
        from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer, pagesize=letter,
            leftMargin=0.7 * inch, rightMargin=0.7 * inch,
            topMargin=0.6 * inch, bottomMargin=0.6 * inch,
        )

        styles = getSampleStyleSheet()
        name_style = ParagraphStyle("Name", parent=styles["Title"], fontSize=20, spaceAfter=2)
        subtitle_style = ParagraphStyle("Subtitle", parent=styles["Normal"], fontSize=10, textColor="#444444")
        heading_style = ParagraphStyle("Heading", parent=styles["Heading2"], fontSize=12, spaceBefore=10, spaceAfter=4, textColor="#1a1a1a")
        body_style = ParagraphStyle("Body", parent=styles["Normal"], fontSize=10, leading=13, spaceAfter=2)
        bullet_style = ParagraphStyle("Bullet", parent=styles["Normal"], fontSize=10, leading=13, leftIndent=14, bulletIndent=4, spaceAfter=1)

        story: list = []
        story.append(Paragraph(self.full_name, name_style))
        story.append(Paragraph(f"{self.title} | {self.email} | {self.phone} | {self.location}", subtitle_style))
        story.append(Spacer(1, 6))
        story.append(HRFlowable(width="100%", thickness=1, color="#888888"))

        def _section(title: str, blocks: list) -> None:
            story.append(Paragraph(title, heading_style))
            for block in blocks:
                story.extend(block)
            story.append(Spacer(1, 2))

        story.append(Paragraph("Professional Summary", heading_style))
        story.append(Paragraph(self.summary, body_style))

        skill_blocks = [[Paragraph(", ".join(self.skills), body_style)]]
        _section("Skills", skill_blocks)

        exp_blocks = []
        for exp in self.experiences:
            exp_blocks.append([
                Paragraph(f"<b>{exp.title}</b> at {exp.company}", body_style),
                Paragraph(f"{exp.start} - {exp.end}", subtitle_style),
            ])
            exp_blocks.append([Paragraph(f"&bull; {b}", bullet_style) for b in exp.bullets])
        _section("Professional Experience", exp_blocks)

        edu_blocks = []
        for edu in self.education:
            gpa = f" | GPA: {edu.gpa}" if edu.gpa else ""
            edu_blocks.append([
                Paragraph(f"{edu.degree} - {edu.institution}", body_style),
                Paragraph(f"{edu.start} - {edu.end}{gpa}", subtitle_style),
            ])
        _section("Education", edu_blocks)

        proj_blocks = []
        for proj in self.projects:
            proj_blocks.append([
                Paragraph(f"<b>{proj.name}</b>", body_style),
                Paragraph(f"Technologies: {', '.join(proj.technologies)}", subtitle_style),
            ])
            proj_blocks.append([Paragraph(f"&bull; {b}", bullet_style) for b in proj.bullets])
        _section("Projects", proj_blocks)

        cert_blocks = [[Paragraph(c, body_style)] for c in self.certifications]
        _section("Certifications", cert_blocks)

        lang_blocks = [[Paragraph(f"{l['name']} - {l['proficiency']}", body_style)] for l in self.languages]
        _section("Languages", lang_blocks)

        doc.build(story)
        return buffer.getvalue()

    def render(self, fmt: str) -> tuple[bytes, str]:
        ext = {"pdf": ".pdf", "docx": ".docx", "txt": ".txt"}.get(fmt, ".txt")
        if fmt == "pdf":
            return self.to_pdf_bytes(), ext
        if fmt == "docx":
            return self.to_docx_bytes(), ext
        return self.to_text().encode("utf-8"), ext


def build_resume_content(
    *,
    rng,
    full_name: str,
    title: str,
    email: str,
    phone: str,
    location: str,
    skills: list[str],
    companies: list[str],
    education_degrees: list[str],
    education_institutions: list[str],
    certifications: list[str],
) -> ResumeContent:
    summary = rng.choice(data_pool.RESUME_SUMMARIES)

    experiences: list[ResumeExperience] = []
    num_experiences = rng.randint(2, 4)
    start_year = rng.randint(2013, 2020)
    for idx in range(num_experiences):
        exp_start = start_year + idx * 2
        exp_end = None
        is_current = idx == num_experiences - 1
        exp_end = "Present" if is_current else start_year + idx * 2 + 1
        bullets = rng.sample(data_pool.EXPERIENCE_BULLETS, k=min(3, len(data_pool.EXPERIENCE_BULLETS)))
        company = rng.choice(companies)
        exp_title = title if idx == 0 else rng.choice(
            [t for t in data_pool.JOB_TITLES if t != title] or data_pool.JOB_TITLES
        )
        experiences.append(
            ResumeExperience(
                title=exp_title,
                company=company,
                start=str(exp_start),
                end=exp_end,
                bullets=bullets,
            )
        )

    education: list[ResumeEducation] = []
    degree = rng.choice(education_degrees)
    institution = rng.choice(education_institutions)
    edu_end_year = rng.randint(2010, 2016)
    education.append(
        ResumeEducation(
            degree=degree,
            institution=institution,
            start=str(edu_end_year - 4),
            end=str(edu_end_year),
            gpa=f"{rng.uniform(3.2, 3.9):.2f}",
        )
    )
    if rng.random() < 0.5:
        degree2 = rng.choice([d for d in education_degrees if d != degree] or education_degrees)
        edu_end_year2 = rng.randint(2016, 2019)
        education.append(
            ResumeEducation(
                degree=degree2,
                institution=rng.choice(education_institutions),
                start=str(edu_end_year2 - 2),
                end=str(edu_end_year2),
            )
        )

    projects: list[ResumeProject] = []
    for proj_name in rng.sample(data_pool.PROJECT_NAMES, k=min(2, len(data_pool.PROJECT_NAMES))):
        proj_techs = rng.sample(skills, k=min(3, len(skills)))
        projects.append(
            ResumeProject(
                name=proj_name,
                technologies=proj_techs,
                bullets=[
                    "Built end-to-end feature with automated tests.",
                    "Improved performance and observability in production.",
                ],
            )
        )

    languages = [
        {"name": "English", "proficiency": "Native"},
        {"name": rng.choice(["Hindi", "Spanish", "French", "German"]), "proficiency": "Fluent"},
    ]

    return ResumeContent(
        full_name=full_name,
        title=title,
        email=email,
        phone=phone,
        location=location,
        summary=summary,
        skills=skills,
        education=education,
        experiences=experiences,
        projects=projects,
        certifications=certifications,
        languages=languages,
    )
